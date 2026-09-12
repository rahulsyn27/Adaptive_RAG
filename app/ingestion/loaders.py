"""Document loaders for PDF, text, Markdown, and DOCX."""

from pathlib import Path

from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document

from app.core.constants import SUPPORTED_DOCUMENT_EXTENSIONS
from app.core.exceptions import EmptyDocumentError, InvalidDocumentTypeError


def _load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    parsed = DocxDocument(str(path))
    paragraphs = [p.text.strip() for p in parsed.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    return [Document(page_content=text, metadata={"source": str(path), "page": None})]


def load_file(path: Path) -> list[Document]:
    """Load a supported file into LangChain documents."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise InvalidDocumentTypeError(path.name)
    if suffix == ".pdf":
        documents = PyPDFLoader(str(path)).load()
    elif suffix == ".docx":
        documents = _load_docx(path)
    else:
        documents = TextLoader(str(path), encoding="utf-8").load()
    nonempty = [doc for doc in documents if doc.page_content and doc.page_content.strip()]
    if not nonempty:
        raise EmptyDocumentError(path.name)
    return nonempty
